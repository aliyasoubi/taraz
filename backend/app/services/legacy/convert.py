"""Old workbook → `StandardBook`.

Per patient and treatment case, rows are replayed in date order against a
running balance:

* a payment becomes a payment line (who paid → which account);
* when the recorded balance is higher than expected, the difference is the
  treatment that was done — a charge line;
* when it is lower without a payment, an adjustment (discount) line, flagged;
* a sold item or visit fee ("visit", "toothbrush" in the balance cell) is a
  charge settled by the same payment; a deposit leaves the patient in credit.

Finally the debtors list, the practice's own current receivables, is taken as
the truth: each listed patient's balance is aligned to it and its cheques are
added as pending cheques, so the debt stays open until they clear.
"""

from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date
from itertools import combinations
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.core.i18n import t
from app.models import ChequeStatus, PaymentMethod, TreatmentKind, TxType
from app.services.legacy.cells import (
    Amount,
    PatientName,
    edit_distance,
    parse_amount,
    parse_date,
    parse_name,
    text_of,
)
from app.services.legacy.destinations import Accounts, Destination
from app.services.workbook.model import Issue, Line, PatientEntry, StandardBook
from app.utils.jalali import JMonth
from app.utils.text import fold

CaseKey = tuple[str, TreatmentKind | None]


@dataclass
class _Row:
    source: str
    order: tuple[int, int]
    name: PatientName
    on: date
    balance: Amount
    payment: Amount
    dest: Destination
    notes: list[str]
    date_problem: str | None


@dataclass
class _Debt:
    source: str
    name: PatientName
    amount: int
    cheque: bool
    due: date | None
    payer: str | None
    note: str | None


@dataclass
class _Case:
    balance: int = 0
    seen: bool = False


@dataclass
class _Converter:
    as_of: date
    # Treat the debtors list as complete: anyone not on it is settled at `as_of`.
    settle_unlisted: bool = False
    accounts: Accounts = field(default_factory=Accounts)
    rows: list[_Row] = field(default_factory=list)
    debts: list[_Debt] = field(default_factory=list)
    book: StandardBook = field(default_factory=StandardBook)
    names: dict[str, Counter[str]] = field(default_factory=lambda: defaultdict(Counter))
    cases: dict[CaseKey, _Case] = field(default_factory=lambda: defaultdict(_Case))

    # ── Reading ──────────────────────────────────────────────────────────────

    def issue(self, source: str, patient: str, text: str) -> str:
        self.book.issues.append(Issue(source, patient, text))
        return text

    def read(self, path: Path) -> None:
        wb = load_workbook(path, data_only=True, read_only=True)
        month_sheets = [ws for ws in wb.worksheets if ws.title in t("legacy.monthSheets")]
        sheets = {ws.title: ws for ws in wb.worksheets}
        bank_cells: list[object] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(min_row=2, values_only=True):
                bank_cells.extend(row[3:7])
        self.accounts.learn_doctors(bank_cells)

        months = sorted((self._sheet_month(ws), ws) for ws in month_sheets if self._sheet_month(ws))
        for index, (month, ws) in enumerate(months):
            self._read_rows(ws, index, month, cols=(0, 1, 2, 3, 4, 5))
        if ortho := sheets.get(t("legacy.orthoSheet")):
            # name, plan total, balance, payment, date, account, notes
            self._read_rows(
                ortho, len(months), None, cols=(0, 2, 3, 4, 5, 6), case=TreatmentKind.ORTHO
            )
        if debtors := sheets.get(t("legacy.debtorsSheet")):
            self._read_debtors(debtors)

    @staticmethod
    def _sheet_month(ws: Worksheet) -> tuple[int, int] | None:
        found: Counter[tuple[int, int]] = Counter()
        for row in ws.iter_rows(min_row=2, max_col=4, values_only=True):
            parsed = parse_date(row[3] if len(row) > 3 else None, None, None)
            if parsed.on and not parsed.problem:
                m = JMonth.of(parsed.on)
                found[(m.year, m.month)] += 1
        return found.most_common(1)[0][0] if found else None

    def _read_rows(
        self,
        ws: Worksheet,
        sheet_index: int,
        month: tuple[int, int] | None,
        cols: tuple[int, int, int, int, int, int],
        case: TreatmentKind | None = None,
    ) -> None:
        name_c, balance_c, payment_c, date_c, bank_c, note_c = cols
        previous: date | None = None
        for line, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            cells = [*row, *([None] * 8)]
            name = parse_name(cells[name_c])
            if name is None:
                continue
            if case is not None:
                name = PatientName(name.display, name.key, case)
            source = f"{ws.title}!{line}"
            parsed = parse_date(cells[date_c], month, previous)
            if parsed.on is None:
                self.issue(source, name.display, t("legacy.issues.skippedRow"))
                continue
            previous = parsed.on
            dest = self.accounts.read(cells[bank_c], cells[note_c])
            self.names[name.key][name.display] += 1
            self.rows.append(
                _Row(
                    source=source,
                    order=(sheet_index, line),
                    name=name,
                    on=parsed.on,
                    balance=parse_amount(cells[balance_c]),
                    payment=parse_amount(cells[payment_c]),
                    dest=dest,
                    notes=dest.notes,
                    date_problem=parsed.problem,
                )
            )

    def _read_debtors(self, ws: Worksheet) -> None:
        cheque_word = fold(t("legacy.cheque"))
        for line, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            name_cell, amount_cell, date_cell, bank_cell, note_cell = [
                *row,
                None,
                None,
                None,
                None,
                None,
            ][:5]
            name = parse_name(name_cell)
            if name is None:
                continue
            source = f"{ws.title}!{line}"
            amount = parse_amount(amount_cell)
            if amount.value is None:
                self.issue(source, name.display, t("legacy.issues.unknownBalance", raw=amount.raw))
                continue
            date_text, bank_text = text_of(date_cell), text_of(bank_cell)
            due = (
                parse_date(date_cell, None, None).on
                if any(c.isdigit() for c in date_text)
                else None
            )
            cheque = fold(bank_text).startswith(cheque_word) or fold(date_text) == cheque_word
            payer = (
                bank_text[len(t("legacy.cheque")) :].strip()
                if fold(bank_text).startswith(cheque_word)
                else None
            )
            note = " — ".join(
                x
                for x in (
                    date_text if due is None and not cheque else "",
                    bank_text if not cheque else "",
                    text_of(note_cell),
                )
                if x
            )
            self.names[name.key][name.display] += 1
            self.debts.append(
                _Debt(source, name, amount.value, cheque, due, payer or None, note or None)
            )

    # ── Replaying the ledger ─────────────────────────────────────────────────

    def _line(self, name: PatientName, **fields: Any) -> Line:
        line = Line(patient=name.key, case=name.case, **fields)
        self.book.lines.append(line)
        return line

    def replay(self) -> None:
        for row in sorted(self.rows, key=lambda r: (r.on, r.order)):
            self._replay(row)

    def _replay(self, row: _Row) -> None:
        name, dest = row.name, row.dest
        cheque_word = fold(t("legacy.cheque"))
        balance_text = fold(row.balance.text or "")
        payment_text = fold(row.payment.text or "")

        if dest.method == PaymentMethod.CHEQUE and not row.payment.value and row.balance.value:
            self.issue(row.source, name.display, t("legacy.issues.chequeSchedule"))
            return
        if balance_text.startswith(cheque_word) and row.payment.value is None:
            self.issue(row.source, name.display, t("legacy.issues.skippedRow"))
            return
        if row.balance.value is None and row.payment.value is None and not row.balance.text:
            return

        state = self.cases[(name.key, name.case)]
        state.seen = True
        paid = row.payment.value or 0
        sold = bool(paid) and any(fold(w) in balance_text for w in t("legacy.salesWords"))
        deposit = fold(t("legacy.deposit")) in balance_text
        known = row.balance.value
        expected = state.balance - paid + (paid if sold else 0)
        delta = known - expected if known is not None else 0
        paid_by_cheque = delta < 0 and cheque_word in payment_text

        # Everything doubtful about the row goes to the issues sheet and onto its lines.
        flags: list[str] = []
        if row.date_problem:
            flags.append(row.date_problem)
        for amount in (row.payment, row.balance):
            if amount.ambiguous:
                flags.append(t("legacy.issues.ambiguousAmount", raw=amount.raw, fixed=amount.value))
            if amount.tiny:
                flags.append(t("legacy.issues.tinyAmount", raw=amount.value))
        if row.payment.value is None and row.payment.text and not paid_by_cheque:
            flags.append(t("legacy.issues.unknownPayment", raw=row.payment.raw))
        if known is None and row.balance.text and not sold and not deposit:
            flags.append(t("legacy.issues.unknownBalance", raw=row.balance.raw))
        if delta < 0 and not paid_by_cheque:
            flags.append(t("legacy.issues.balanceDrop", amount=-delta))
        for text in flags:
            self.issue(row.source, name.display, text)
        common: dict[str, Any] = {
            "on": row.on,
            "source": row.source,
            "doctor": dest.doctor,
            "review": "\u061b ".join(flags) or None,
        }

        if sold:  # an item or visit fee, paid on the spot
            self._line(
                name, type=TxType.CHARGE, amount=paid, description=row.balance.text, **common
            )
        if delta > 0:  # treatment done since the previous row
            self._line(
                name, type=TxType.CHARGE, amount=delta, description=t("legacy.chargeNote"), **common
            )
        if paid:
            self.accounts.use(dest)
            cheque = dest.method == PaymentMethod.CHEQUE
            self._line(
                name,
                type=TxType.PAYMENT,
                amount=paid,
                payer=dest.payer,
                account=dest.account,
                method=dest.method,
                cheque_due=row.on if cheque else None,
                cheque_status=ChequeStatus.CLEARED if cheque else None,
                description=t("legacy.deposit") if deposit else (" — ".join(row.notes) or None),
                **common,
            )
        if paid_by_cheque:  # "paid by cheque" with no amount: the balance drop is the payment
            self.accounts.use(dest)
            self._line(
                name,
                type=TxType.PAYMENT,
                amount=-delta,
                account=dest.account,
                method=PaymentMethod.CHEQUE,
                cheque_due=row.on,
                cheque_status=ChequeStatus.CLEARED,
                description=row.payment.text,
                **common,
            )
        elif delta < 0:
            self._line(
                name,
                type=TxType.DISCOUNT,
                amount=-delta,
                description=t("legacy.adjustNote"),
                **common,
            )
        state.balance = known if known is not None else expected

    def apply_debtors(self) -> None:
        listed: dict[CaseKey, list[_Debt]] = defaultdict(list)
        for debt in self.debts:
            listed[(debt.name.key, debt.name.case)].append(debt)

        for key, debts in listed.items():
            state = self.cases[key]
            name = debts[0].name
            target = sum(d.amount for d in debts)  # cheques are debt until they clear
            delta = target - state.balance
            if delta:
                if state.seen:
                    review = self.issue(
                        debts[0].source,
                        name.display,
                        t("legacy.issues.debtorsMismatch", computed=state.balance, listed=target),
                    )
                    note = t("legacy.debtorsAdjustNote")
                else:
                    review, note = None, t("legacy.openingNote")
                self._line(
                    name,
                    on=self.as_of,
                    type=TxType.CHARGE if delta > 0 else TxType.DISCOUNT,
                    amount=abs(delta),
                    description=" — ".join(
                        x for x in (note, *(d.note for d in debts if d.note and not d.cheque)) if x
                    ),
                    source=debts[0].source,
                    review=review,
                )
                state.balance = target
            cheques = [d for d in debts if d.cheque]
            if cheques and any(not d.cheque for d in debts):
                self.issue(debts[0].source, name.display, t("legacy.issues.chequesAndBalance"))
            for d in cheques:
                review = (
                    None
                    if d.due
                    else self.issue(d.source, name.display, t("legacy.issues.chequeUnknownDue"))
                )
                self._line(
                    name,
                    on=self.as_of,
                    type=TxType.PAYMENT,
                    amount=d.amount,
                    method=PaymentMethod.CHEQUE,
                    cheque_due=d.due,
                    cheque_status=ChequeStatus.PENDING,
                    payer=d.payer,
                    description=t("legacy.chequeNote"),
                    source=d.source,
                    review=review,
                )

        for (pkey, case), state in self.cases.items():
            if state.balance <= 0 or (pkey, case) in listed:
                continue
            name = PatientName(self._display(pkey), pkey, case)
            if not self.settle_unlisted:
                self.issue("", name.display, t("legacy.issues.notInDebtors", amount=state.balance))
                continue
            review = self.issue(
                "", name.display, t("legacy.issues.settledUnlisted", amount=state.balance)
            )
            self._line(
                name,
                on=self.as_of,
                type=TxType.DISCOUNT,
                amount=state.balance,
                description=t("legacy.settleNote"),
                review=review,
            )
            state.balance = 0

    # ── Output ───────────────────────────────────────────────────────────────

    def _display(self, key: str) -> str:
        return self.names[key].most_common(1)[0][0] if self.names.get(key) else key

    def finish(self) -> StandardBook:
        book = self.book
        for line in book.lines:
            line.patient = self._display(line.patient)

        entries: dict[str, PatientEntry] = {}
        for key, spellings in self.names.items():
            display = self._display(key)
            entries[key] = PatientEntry(
                name=display, aliases={s for s in spellings if s != display}
            )
        keys = sorted(entries)
        for a, b in combinations(keys, 2):
            if min(len(a), len(b)) >= 6 and edit_distance(a, b, limit=2) <= 1:
                entries[a].possible_duplicates.append(entries[b].name)
                entries[b].possible_duplicates.append(entries[a].name)
                self.issue(
                    "", entries[a].name, t("legacy.issues.possibleDuplicate", other=entries[b].name)
                )

        by_name = {e.name: e for e in entries.values()}
        for line in book.lines:
            entry = by_name[line.patient]
            if line.cheque_status == ChequeStatus.PENDING:
                entry.pending_cheques += line.amount
            elif line.cheque_status == ChequeStatus.BOUNCED:
                continue
            elif line.type in (TxType.CHARGE, TxType.REFUND):
                entry.balance += line.amount
            else:
                entry.balance -= line.amount

        book.patients = sorted(entries.values(), key=lambda e: (-e.balance, e.name))
        book.accounts = self.accounts.entries
        book.issues.sort(key=lambda i: (i.source == "", i.source, i.patient))
        return book


def convert_legacy(path: Path, as_of: date, *, settle_unlisted: bool = False) -> StandardBook:
    """Read the old workbook and rebuild it as a reviewable standard book."""
    conv = _Converter(as_of=as_of, settle_unlisted=settle_unlisted)
    conv.read(path)
    conv.replay()
    conv.apply_debtors()
    return conv.finish()
