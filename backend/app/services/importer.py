"""One-time import of the old spreadsheet's debtors sheet as opening balances.

Each row becomes a charge ("balance carried over from Excel") on the patient's
ledger. A row that is a cheque (the account column starts with the word for
cheque) additionally becomes a pending cheque payment for the same amount, so
the debt stays open until the cheque clears — exactly how the sheet treated it.

Names are matched loosely (letter variants, dots, spacing). A trailing `imp`
(or the Persian word for implant / ortho / laminate) moves the row onto that
treatment case of the same patient, as the sheet did with separate "… imp" rows.

The sheet name and every Persian word matched here come from the locale file
(``app/core/i18n/fa.json``, section ``import``).
"""

import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.i18n import t
from app.models import (
    ChequeStatus,
    Patient,
    PaymentMethod,
    Transaction,
    Treatment,
    TreatmentKind,
    TxType,
)
from app.utils.jalali import parse_jalali
from app.utils.text import fold, tidy

SHEET: str = t("import.sheet")
CHEQUE: str = t("import.cheque")


def _suffix_pattern(words: list[str]) -> re.Pattern[str]:
    # Longest first, so "ortho…" wins over its own prefix.
    alternatives = "|".join(re.escape(w) for w in sorted(words, key=len, reverse=True))
    return re.compile(rf"\s*({alternatives})\s*$", re.IGNORECASE)


_SUFFIXES: list[tuple[re.Pattern[str], TreatmentKind, str]] = [
    (_suffix_pattern(words), TreatmentKind(kind), t(f"treatment.{kind}"))
    for kind, words in t("import.suffixes").items()
]


@dataclass
class ParsedRow:
    line: int
    name: str
    amount: int
    kind: TreatmentKind | None
    case_title: str | None
    cheque: bool
    due_on: date | None
    payer: str | None
    note: str | None


@dataclass
class ImportResult:
    rows: list[ParsedRow] = field(default_factory=list)
    skipped: list[tuple[int, str]] = field(default_factory=list)
    patients_created: int = 0
    patients_matched: int = 0
    charges: int = 0
    cheques: int = 0
    # Patients with both a plain balance row and cheque rows: the sheet does not
    # say whether the cheques cover that balance or are on top of it.
    review_lines: list[int] = field(default_factory=list)


def _amount(value: object) -> int | None:
    if isinstance(value, int | float) and value > 0:
        return int(value)
    return None


def _split_name(raw: str) -> tuple[str, TreatmentKind | None, str | None]:
    name = tidy(raw.replace(".", " "))  # "First. Last" → "First Last"
    for pattern, kind, title in _SUFFIXES:
        if pattern.search(name):
            return tidy(pattern.sub("", name)), kind, title
    return name, None, None


def parse(path: Path) -> ImportResult:
    wb = load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        raise ValueError(f"sheet {SHEET!r} not found; sheets: {', '.join(wb.sheetnames)}")
    result = ImportResult()
    for line, row in enumerate(wb[SHEET].iter_rows(min_row=2, values_only=True), start=2):
        cells = [*row, None, None, None, None, None][:5]
        raw_name, raw_amount, raw_date, raw_account, raw_note = cells
        if not raw_name or not str(raw_name).strip():
            continue
        amount = _amount(raw_amount)
        if amount is None:
            result.skipped.append((line, f"amount is not a number: {raw_amount!r}"))
            continue
        name, kind, case_title = _split_name(str(raw_name))
        date_text = str(raw_date).strip() if raw_date is not None else ""
        account_text = tidy(str(raw_account)) if raw_account else ""
        due_on = parse_jalali(date_text) if date_text else None

        if date_text and due_on is None and kind is None:
            # A word in the date column, e.g. "implant" or "cheque".
            for pattern, k, title in _SUFFIXES:
                if pattern.search(" " + date_text):
                    kind, case_title = k, title
        cheque = account_text.startswith(CHEQUE) or date_text == CHEQUE
        payer = tidy(account_text[len(CHEQUE) :]) if account_text.startswith(CHEQUE) else None

        notes = [n for n in (str(raw_note).strip() if raw_note else "",) if n]
        if not cheque and account_text:
            notes.append(account_text)
        if not cheque and due_on:
            notes.append(t("import.dueNote", date=date_text))
        if date_text and due_on is None and date_text != CHEQUE and kind is None:
            notes.append(date_text)
        result.rows.append(
            ParsedRow(
                line=line,
                name=name,
                amount=amount,
                kind=kind,
                case_title=case_title,
                cheque=cheque,
                due_on=due_on if cheque else None,
                payer=payer or None,
                note=" — ".join(notes) or None,
            )
        )

    by_patient: dict[str, list[ParsedRow]] = {}
    for r in result.rows:
        by_patient.setdefault(fold(r.name), []).append(r)
    for rows in by_patient.values():
        if any(r.cheque for r in rows) and any(not r.cheque for r in rows):
            result.review_lines.extend(r.line for r in rows)
    return result


def run(db: Session, path: Path, as_of: date, *, dry_run: bool = False) -> ImportResult:
    result = parse(path)
    patients: dict[str, Patient] = {p.name_search: p for p in db.scalars(select(Patient)).all()}
    preexisting = set(patients)
    cases: dict[tuple[int, TreatmentKind], Treatment] = {}

    for r in result.rows:
        key = fold(r.name)
        patient = patients.get(key)
        if patient is None:
            patient = Patient(name=r.name, name_search=key)
            db.add(patient)
            db.flush()
            patients[key] = patient
            result.patients_created += 1

        treatment_id = None
        if r.kind is not None:
            ck = (patient.id, r.kind)
            if ck not in cases:
                existing = db.scalars(
                    select(Treatment).where(
                        Treatment.patient_id == patient.id, Treatment.kind == r.kind
                    )
                ).first()
                cases[ck] = existing or Treatment(
                    patient_id=patient.id, title=r.case_title or r.kind.value, kind=r.kind
                )
                db.add(cases[ck])
                db.flush()
            treatment_id = cases[ck].id

        db.add(
            Transaction(
                patient_id=patient.id,
                treatment_id=treatment_id,
                type=TxType.CHARGE,
                amount=r.amount,
                occurred_on=as_of,
                description=t("import.openingNote", line=r.line)
                + (f" — {r.note}" if r.note else ""),
            )
        )
        result.charges += 1
        if r.cheque:
            db.add(
                Transaction(
                    patient_id=patient.id,
                    treatment_id=treatment_id,
                    type=TxType.PAYMENT,
                    amount=r.amount,
                    occurred_on=as_of,
                    method=PaymentMethod.CHEQUE,
                    cheque_status=ChequeStatus.PENDING,
                    cheque_due_on=r.due_on,
                    payer_name=r.payer,
                    description=t("import.chequeNote", line=r.line),
                )
            )
            result.cheques += 1

    result.patients_matched = len({fold(r.name) for r in result.rows} & preexisting)
    if dry_run:
        db.rollback()
    else:
        db.commit()
    return result


def has_transactions(db: Session) -> bool:
    return bool(db.scalar(select(func.count(Transaction.id))))
