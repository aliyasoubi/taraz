"""`StandardBook` → Excel, laid out for review: one sheet each for transactions,
patients, accounts and open questions, with dropdowns where a value is a choice."""

from collections.abc import Iterable, Sequence
from enum import StrEnum
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from app.core.i18n import t
from app.models import AccountKind, ChequeStatus, PaymentMethod, TreatmentKind, TxType
from app.services.spreadsheet.cells import jalali_text
from app.services.workbook.model import StandardBook

HEADER_FILL = PatternFill("solid", fgColor="4C1D95")
REVIEW_FILL = PatternFill("solid", fgColor="FEF3C7")
MONEY = "#,##0"


def label(group: str, value: StrEnum | None) -> str:
    return str(t(f"{group}.{value.value}")) if value is not None else ""


def labels(group: str, enum: type[StrEnum]) -> list[str]:
    return [label(group, member) for member in enum]


def _sheet(
    wb: Workbook, title: str, headers: Sequence[str], rows: Iterable[Sequence[Any]]
) -> Worksheet:
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.append(list(headers))
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append(list(row))
    for idx, header in enumerate(headers, 1):
        letter = get_column_letter(idx)
        width = max([len(str(header))] + [len(str(c.value or "")) for c in ws[letter][1:200]])
        ws.column_dimensions[letter].width = min(max(10, width + 2), 60)
    ws.freeze_panes = "A2"
    return ws


def _dropdown(ws: Worksheet, column: int, options: list[str]) -> None:
    letter = get_column_letter(column)
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    dv.add(f"{letter}2:{letter}{max(ws.max_row, 2) + 500}")
    ws.add_data_validation(dv)


def write_book(book: StandardBook) -> bytes:
    wb = Workbook()
    names = t("workbook.sheets")
    general = t("workbook.general")

    guide = _sheet(wb, names["guide"], [names["guide"]], ([line] for line in t("workbook.guide")))
    guide.column_dimensions["A"].width = 110

    tx = _sheet(
        wb,
        names["transactions"],
        t("workbook.transactions"),
        (
            (
                jalali_text(line.on),
                line.patient,
                label("treatment", line.case) or general,
                label("tx", line.type),
                line.amount,
                line.payer or "",
                line.account or "",
                label("method", line.method),
                jalali_text(line.cheque_due) if line.cheque_due else "",
                label("cheque", line.cheque_status),
                line.doctor or "",
                line.description or "",
                line.source,
                line.review or "",
            )
            for line in book.lines
        ),
    )
    for (cell,) in tx.iter_rows(min_row=2, min_col=5, max_col=5):
        cell.number_format = MONEY
    for row in tx.iter_rows(min_row=2):
        if row[-1].value:
            for cell in row:
                cell.fill = REVIEW_FILL
    _dropdown(tx, 3, labels("treatment", TreatmentKind))
    _dropdown(tx, 4, labels("tx", TxType))
    _dropdown(tx, 8, labels("method", PaymentMethod))
    _dropdown(tx, 10, labels("cheque", ChequeStatus))

    patients = _sheet(
        wb,
        names["patients"],
        t("workbook.patients"),
        (
            (
                p.name,
                " | ".join(sorted(p.aliases)),
                " | ".join(p.possible_duplicates),
                p.merge_into or "",
                p.balance,
                p.pending_cheques,
            )
            for p in book.patients
        ),
    )
    for column in (5, 6):
        for (cell,) in patients.iter_rows(min_row=2, min_col=column, max_col=column):
            cell.number_format = MONEY

    accounts = _sheet(
        wb,
        names["accounts"],
        t("workbook.accounts"),
        (
            (
                a.title,
                a.bank or "",
                a.owner or "",
                label("accountKind", a.kind),
                " | ".join(sorted(a.aliases)),
                a.uses,
            )
            for a in book.accounts
        ),
    )
    _dropdown(accounts, 4, labels("accountKind", AccountKind))

    _sheet(
        wb,
        names["issues"],
        t("workbook.issues"),
        ((i.source, i.patient, i.text) for i in book.issues),
    )

    del wb["Sheet"]
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
