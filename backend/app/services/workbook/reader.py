"""Excel → `StandardBook`, validating every row of an edited standard workbook."""

from dataclasses import dataclass, field
from datetime import date
from enum import StrEnum
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.core.i18n import t
from app.models import AccountKind, ChequeStatus, PaymentMethod, TreatmentKind, TxType
from app.services.legacy.cells import parse_amount, parse_date, text_of
from app.services.workbook.model import AccountEntry, Line, PatientEntry, StandardBook
from app.services.workbook.writer import label
from app.utils.text import fold

# The same limits as the database columns and the API, so an imported row can
# later be edited in the app.
_LIMITS = {"patient": 120, "payer": 120, "account": 80, "doctor": 80, "description": 300}


@dataclass
class ReadResult:
    book: StandardBook
    errors: list[str] = field(default_factory=list)


def _choice[E: StrEnum](group: str, enum: type[E], text: str) -> E | None:
    by_label = {fold(label(group, m)): m for m in enum}
    return by_label.get(fold(text)) if text else None


def _columns(header: tuple[Any, ...], expected: list[str], sheet: str) -> dict[str, int]:
    found = {fold(text_of(h)): i for i, h in enumerate(header)}
    missing = [h for h in expected if fold(h) not in found]
    if missing:
        raise ValueError(f"{sheet}: missing columns {missing}")
    return {h: found[fold(h)] for h in expected}


def read_book(path: Path) -> ReadResult:
    wb = load_workbook(path, data_only=True, read_only=True)
    names = t("workbook.sheets")
    result = ReadResult(StandardBook())
    _read_accounts(wb[names["accounts"]], result)
    _read_patients(wb[names["patients"]], result)
    _read_lines(wb[names["transactions"]], result)
    return result


def _rows(ws: Any, headers: list[str]) -> Any:
    rows = ws.iter_rows(values_only=True)
    cols = _columns(next(rows), headers, ws.title)
    for number, row in enumerate(rows, start=2):
        values = {h: row[i] if i < len(row) else None for h, i in cols.items()}
        if any(text_of(v) for v in values.values()):
            yield number, values


def _read_accounts(ws: Any, result: ReadResult) -> None:
    title, bank, owner, kind, _aliases, _uses = t("workbook.accounts")
    for number, row in _rows(ws, t("workbook.accounts")):
        account_kind = _choice("accountKind", AccountKind, text_of(row[kind])) or AccountKind.OTHER
        if not text_of(row[title]):
            result.errors.append(f"{ws.title}!{number}: empty title")
            continue
        result.book.accounts.append(
            AccountEntry(
                title=text_of(row[title]),
                kind=account_kind,
                bank=text_of(row[bank]) or None,
                owner=text_of(row[owner]) or None,
            )
        )


def _read_patients(ws: Any, result: ReadResult) -> None:
    name, _aliases, _dupes, merge_into, _balance, _pending = t("workbook.patients")
    for _number, row in _rows(ws, t("workbook.patients")):
        result.book.patients.append(
            PatientEntry(name=text_of(row[name]), merge_into=text_of(row[merge_into]) or None)
        )


def _read_lines(ws: Any, result: ReadResult) -> None:
    (
        on,
        patient,
        case,
        kind,
        amount,
        payer,
        account,
        method,
        due,
        status,
        doctor,
        note,
        source,
        review,
    ) = t("workbook.transactions")
    general = fold(t("workbook.general"))
    for number, row in _rows(ws, t("workbook.transactions")):
        where = f"{ws.title}!{number}"
        problems: list[str] = []
        parsed_on = parse_date(row[on], None, None)
        tx_type = _choice("tx", TxType, text_of(row[kind]))
        parsed_amount = parse_amount(row[amount])
        value = parsed_amount.value
        due_on: date | None = None
        if text_of(row[due]):
            due_on = parse_date(row[due], None, None).on
            if due_on is None:
                problems.append("cheque due date")
        case_text = text_of(row[case])
        case_kind = (
            None
            if fold(case_text) in ("", general)
            else _choice("treatment", TreatmentKind, case_text)
        )
        if case_text and fold(case_text) != general and case_kind is None:
            problems.append("case")
        method_text, status_text = text_of(row[method]), text_of(row[status])
        pay_method = _choice("method", PaymentMethod, method_text)
        cheque_status = _choice("cheque", ChequeStatus, status_text)
        if method_text and pay_method is None:
            problems.append("method")
        if status_text and cheque_status is None:
            problems.append("cheque status")
        if parsed_on.on is None or parsed_on.problem:
            problems.append("date")
        if tx_type is None:
            problems.append("type")
        if not value or parsed_amount.negative:
            problems.append("amount")
        texts = {
            "patient": text_of(row[patient]),
            "payer": text_of(row[payer]),
            "account": text_of(row[account]),
            "doctor": text_of(row[doctor]),
            # The importer appends the doctor to the description.
            "description": " — ".join(x for x in (text_of(row[note]), text_of(row[doctor])) if x),
        }
        problems.extend(
            f"{field} too long" for field, text in texts.items() if len(text) > _LIMITS[field]
        )
        if not text_of(row[patient]):
            problems.append("patient")
        if problems:
            result.errors.append(f"{where}: invalid {', '.join(problems)}")
            continue
        assert parsed_on.on is not None and tx_type is not None and value
        result.book.lines.append(
            Line(
                on=parsed_on.on,
                patient=text_of(row[patient]),
                type=tx_type,
                amount=value,
                case=case_kind,
                payer=text_of(row[payer]) or None,
                account=text_of(row[account]) or None,
                method=pay_method,
                cheque_due=due_on,
                cheque_status=cheque_status,
                doctor=text_of(row[doctor]) or None,
                description=text_of(row[note]) or None,
                source=text_of(row[source]) or where,
                review=text_of(row[review]) or None,
            )
        )
