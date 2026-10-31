"""One month sheet of the practice's workbook → `StandardBook`.

A month sheet (Aban, then Azar …) has one row per visit or payment: patient,
service, balance after the row, payment, date, the account it went to, notes.
Columns are found by their headers, so an added "payer" or "doctor" column is
read too. Only that one sheet is read.

Each patient's rows are replayed in date order against one running balance,
the sheet's own "balance" column. It starts from what the patient already owes
in the app at the start of the month (`opening`), or from zero:

* a payment becomes a payment line (who paid → which account);
* when the recorded balance is higher than expected, the difference is the
  treatment that was done (or, on a patient's first row, what was owed before)
  — a charge line, named after the service;
* when it is lower without a payment, an adjustment (discount) line, flagged;
* an item or visit fee (toothbrush, visit) is a charge settled by the same
  payment and leaves the balance alone; a deposit leaves the patient in credit.
"""

from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date
from itertools import combinations
from pathlib import Path
from typing import IO, Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.core.errors import InvalidError
from app.core.i18n import t
from app.models import ChequeStatus, PaymentMethod, TxType
from app.services.spreadsheet.cells import (
    Amount,
    PatientName,
    edit_distance,
    parse_amount,
    parse_date,
    parse_name,
    text_of,
)
from app.services.spreadsheet.destinations import Accounts, Destination
from app.services.workbook.model import Issue, Line, PatientEntry, StandardBook
from app.utils.jalali import JMonth
from app.utils.text import fold, tidy

# Where each column sits when a sheet's header does not say.
_DEFAULT_COLUMNS = {
    "name": 0,
    "service": 1,
    "balance": 2,
    "payment": 3,
    "date": 4,
    "account": 5,
    "notes": 6,
}


@dataclass
class _Row:
    source: str
    line: int
    name: PatientName
    on: date
    service: str
    balance: Amount
    payment: Amount
    dest: Destination
    doctor: str | None
    payer: str | None
    date_problem: str | None


def _columns(ws: Worksheet) -> dict[str, int]:
    """Field → column index, from the header row; unknown headers keep the default layout."""
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    found: dict[str, int] = {}
    for index, cell in enumerate(header):
        text = fold(text_of(cell))
        if not text:
            continue
        for field_name, words in t("spreadsheet.headers").items():
            if field_name not in found and any(text == fold(w) for w in words):
                found[field_name] = index
    return {**_DEFAULT_COLUMNS, **found}


def _sheet_month(ws: Worksheet) -> JMonth | None:
    """The month most of the sheet's dates fall in."""
    date_col = _columns(ws)["date"]
    found: Counter[tuple[int, int]] = Counter()
    for row in ws.iter_rows(min_row=2, values_only=True):
        parsed = parse_date(row[date_col] if len(row) > date_col else None, None, None)
        if parsed.on and not parsed.problem:
            m = JMonth.of(parsed.on)
            found[(m.year, m.month)] += 1
    if not found:
        return None
    year, month = found.most_common(1)[0][0]
    return JMonth(year, month)


def pick_month_sheet(path: Path | IO[bytes], month: JMonth | None) -> tuple[Worksheet, JMonth]:
    """The sheet for `month`, or the latest month sheet in the workbook."""
    wb = load_workbook(path, data_only=True, read_only=True)
    dated = [(m, ws) for ws in wb.worksheets if (m := _sheet_month(ws))]
    if month is not None:
        dated = [(m, ws) for m, ws in dated if m == month]
    if not dated:
        raise InvalidError("month_sheet_not_found")
    found, ws = max(dated, key=lambda pair: (pair[0].year, pair[0].month))
    return ws, found


@dataclass
class _Converter:
    month: JMonth
    # name key → what the patient owes at the start of the month
    opening: dict[str, int] = field(default_factory=dict)
    accounts: Accounts = field(default_factory=Accounts)
    rows: list[_Row] = field(default_factory=list)
    book: StandardBook = field(default_factory=StandardBook)
    names: dict[str, Counter[str]] = field(default_factory=lambda: defaultdict(Counter))
    balances: dict[str, int] = field(default_factory=dict)

    # ── Reading ──────────────────────────────────────────────────────────────

    def issue(self, source: str, patient: str, text: str) -> str:
        self.book.issues.append(Issue(source, patient, text))
        return text

    def read(self, ws: Worksheet) -> None:
        cols = _columns(ws)
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        # Doctors are learnt from every text cell first, so "Emami" alone is known to be one.
        self.accounts.learn_doctors(cell for row in rows for cell in row[1:])
        previous: date | None = None
        for line, row in enumerate(rows, start=2):
            cells = [*row, *([None] * 12)]
            cell = {field_name: cells[index] for field_name, index in cols.items()}
            name = parse_name(cell.get("name"))
            if name is None:
                continue
            source = f"{ws.title}!{line}"
            parsed = parse_date(cell.get("date"), (self.month.year, self.month.month), previous)
            if parsed.on is None:
                self.issue(source, name.display, t("spreadsheet.issues.skippedRow"))
                continue
            previous = parsed.on
            dest = self.accounts.read(cell.get("account"), cell.get("notes"))
            self.names[name.key][name.display] += 1
            self.rows.append(
                _Row(
                    source=source,
                    line=line,
                    name=PatientName(name.display, name.key, None),
                    on=parsed.on,
                    service=tidy(text_of(cell.get("service"))),
                    balance=parse_amount(cell.get("balance")),
                    payment=parse_amount(cell.get("payment")),
                    dest=dest,
                    doctor=tidy(text_of(cell.get("doctor"))) or dest.doctor,
                    payer=tidy(text_of(cell.get("payer"))) or dest.payer,
                    date_problem=parsed.problem,
                )
            )

    # ── Replaying the ledger ─────────────────────────────────────────────────

    def _line(self, name: PatientName, **fields: Any) -> Line:
        line = Line(patient=name.key, case=None, **fields)
        self.book.lines.append(line)
        return line

    def replay(self) -> None:
        for row in sorted(self.rows, key=lambda r: (r.on, r.line)):
            self._replay(row)

    def _replay(self, row: _Row) -> None:
        name, dest = row.name, row.dest
        cheque_word = fold(t("spreadsheet.cheque"))
        balance_text = fold(row.balance.text or "")
        payment_text = fold(row.payment.text or "")
        # What the row was for: the service column, or words written in the balance cell.
        about = f"{fold(row.service)} {balance_text}"

        if dest.method == PaymentMethod.CHEQUE and not row.payment.value and row.balance.value:
            self.issue(row.source, name.display, t("spreadsheet.issues.chequeSchedule"))
            return
        if balance_text.startswith(cheque_word) and row.payment.value is None:
            self.issue(row.source, name.display, t("spreadsheet.issues.skippedRow"))
            return
        if row.balance.value is None and row.payment.value is None and not row.balance.text:
            return

        first = name.key not in self.balances
        if first:
            self.balances[name.key] = self.opening.get(name.key, 0)
        before = self.balances[name.key]
        paid = row.payment.value or 0
        deposit = fold(t("spreadsheet.deposit")) in about
        sold = (
            bool(paid)
            and not deposit
            and any(fold(w) in about for w in t("spreadsheet.salesWords"))
        )
        known = row.balance.value
        # A sale stands on its own, and a deposit written with a zero balance leaves
        # the patient in credit: neither says anything about what is still owed.
        if sold or (deposit and not known):
            known = None
        expected = before - paid + (paid if sold else 0)
        delta = known - expected if known is not None else 0
        paid_by_cheque = delta < 0 and cheque_word in payment_text

        # Everything doubtful about the row goes to the issues sheet and onto its lines.
        flags: list[str] = []
        if row.date_problem:
            flags.append(row.date_problem)
        for amount in (row.payment, row.balance):
            if amount.ambiguous:
                flags.append(
                    t("spreadsheet.issues.ambiguousAmount", raw=amount.raw, fixed=amount.value)
                )
            if amount.tiny:
                flags.append(t("spreadsheet.issues.tinyAmount", raw=amount.value))
            if amount.negative:
                flags.append(
                    t(
                        "spreadsheet.issues.negativeAmount",
                        raw=-(amount.value or 0),
                        fixed=amount.value,
                    )
                )
        if row.payment.value is None and row.payment.text and not paid_by_cheque:
            flags.append(t("spreadsheet.issues.unknownPayment", raw=row.payment.raw))
        if row.balance.value is None and row.balance.text and not sold and not deposit:
            flags.append(t("spreadsheet.issues.unknownBalance", raw=row.balance.raw))
        if delta < 0 and not paid_by_cheque:
            flags.append(t("spreadsheet.issues.balanceDrop", amount=-delta))
        if (paid and dest.method == PaymentMethod.CHEQUE) or paid_by_cheque:
            flags.append(t("spreadsheet.issues.chequeAssumedCleared"))
        for text in flags:
            self.issue(row.source, name.display, text)
        common: dict[str, Any] = {
            "on": row.on,
            "source": row.source,
            "doctor": row.doctor,
            "review": "\u061b ".join(flags) or None,
        }
        service = row.service or row.balance.text or None
        notes = " — ".join(row.dest.notes) or None

        if sold:  # an item or visit fee, paid on the spot
            self._line(name, type=TxType.CHARGE, amount=paid, description=service, **common)
        if delta > 0:
            if service and not deposit:
                what = service
            elif first and name.key not in self.opening:
                # The month's first row also carries whatever was owed before it.
                what = t("spreadsheet.openingNote")
            else:
                what = t("spreadsheet.chargeNote")
            self._line(name, type=TxType.CHARGE, amount=delta, description=what, **common)
        if paid:
            self.accounts.use(dest)
            cheque = dest.method == PaymentMethod.CHEQUE
            self._line(
                name,
                type=TxType.PAYMENT,
                amount=paid,
                payer=row.payer,
                account=dest.account,
                method=dest.method,
                cheque_due=row.on if cheque else None,
                cheque_status=ChequeStatus.CLEARED if cheque else None,
                description=" — ".join(x for x in (service, notes) if x) or None,
                **common,
            )
        if paid_by_cheque:  # "paid by cheque" with no amount: the balance drop is the payment
            self.accounts.use(dest)
            self._line(
                name,
                type=TxType.PAYMENT,
                amount=-delta,
                payer=row.payer,
                account=dest.account,
                method=PaymentMethod.CHEQUE,
                cheque_due=row.on,
                cheque_status=ChequeStatus.CLEARED,
                description=row.payment.text,
                **common,
            )
        elif delta < 0:
            self._line(
                name,
                type=TxType.DISCOUNT,
                amount=-delta,
                description=t("spreadsheet.adjustNote"),
                **common,
            )
        self.balances[name.key] = known if known is not None else expected

    # ── Output ───────────────────────────────────────────────────────────────

    def _display(self, key: str) -> str:
        return self.names[key].most_common(1)[0][0] if self.names.get(key) else key

    def finish(self) -> StandardBook:
        book = self.book
        for line in book.lines:
            line.patient = self._display(line.patient)

        entries: dict[str, PatientEntry] = {}
        for key, spellings in self.names.items():
            display = self._display(key)
            entries[key] = PatientEntry(
                name=display,
                aliases={s for s in spellings if s != display},
                balance=self.opening.get(key, 0),
            )
        keys = sorted(entries)
        for a, b in combinations(keys, 2):
            if min(len(a), len(b)) >= 6 and edit_distance(a, b, limit=2) <= 1:
                entries[a].possible_duplicates.append(entries[b].name)
                entries[b].possible_duplicates.append(entries[a].name)
                self.issue(
                    "",
                    entries[a].name,
                    t("spreadsheet.issues.possibleDuplicate", other=entries[b].name),
                )

        by_name = {e.name: e for e in entries.values()}
        for line in book.lines:
            entry = by_name[line.patient]
            if line.cheque_status == ChequeStatus.PENDING:
                entry.pending_cheques += line.amount
            elif line.cheque_status == ChequeStatus.BOUNCED:
                continue
            elif line.type in (TxType.CHARGE, TxType.REFUND):
                entry.balance += line.amount
            else:
                entry.balance -= line.amount

        book.patients = sorted(entries.values(), key=lambda e: (-e.balance, e.name))
        book.accounts = self.accounts.entries
        book.issues.sort(key=lambda i: (i.source == "", i.source, i.patient))
        return book


def convert_month(
    path: Path | IO[bytes],
    month: JMonth | None = None,
    *,
    opening: dict[str, int] | None = None,
) -> tuple[StandardBook, JMonth]:
    """Read one month sheet (the latest, unless `month` is given) as a reviewable book.

    `opening` maps a patient's name key to what they owe at the start of the
    month, so a month imports on top of the ones already in the app.
    """
    ws, found = pick_month_sheet(path, month)
    conv = _Converter(month=found, opening=opening or {})
    conv.read(ws)
    conv.replay()
    return conv.finish(), found
