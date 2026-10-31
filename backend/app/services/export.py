"""Excel exports: the month sheet, the month-end statement, and a full backup of every row.

Sheet names, headers and labels come from the locale file (``app/core/i18n/fa.json``).
"""

from collections.abc import Iterable, Sequence
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.i18n import t
from app.models import Account, Patient, Transaction, Treatment, TxType
from app.schemas.reports import MonthlyReport, MonthSheet
from app.services import ledger
from app.services.transactions import tx_select
from app.utils.jalali import to_jalali_str

_HEADER_FILL = PatternFill("solid", fgColor="4C1D95")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_MONEY = "#,##0"


def _label(group: str, value: str | None) -> str:
    """Display label of an enum value, e.g. ``_label("method", "pos")``."""
    return str(t(f"{group}.{value}")) if value else ""


def _sheet(
    wb: Workbook,
    title: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[Any]],
    money_cols: Iterable[int] = (),
) -> Worksheet:
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.append(list(headers))
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        ws.append(list(row))
    for col in money_cols:
        for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            cell.number_format = _MONEY
    for idx, header in enumerate(headers, 1):
        width = max(
            [len(str(header))] + [len(str(c.value or "")) for c in ws[get_column_letter(idx)]]
        )
        ws.column_dimensions[get_column_letter(idx)].width = min(max(10, width + 3), 48)
    ws.freeze_panes = "A2"
    return ws


def _to_bytes(wb: Workbook) -> bytes:
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _tx_rows(txs: Iterable[Transaction]) -> Iterable[list[Any]]:
    for tx in txs:
        yield [
            tx.id,
            to_jalali_str(tx.occurred_on),
            tx.patient.name,
            tx.treatment.title if tx.treatment else "",
            _label("tx", tx.type),
            tx.amount,
            tx.payer_name or (t("common.self") if tx.type == TxType.PAYMENT else ""),
            tx.account.title if tx.account else "",
            _label("method", tx.method),
            tx.reference or "",
            to_jalali_str(tx.cheque_due_on),
            _label("cheque", tx.cheque_status),
            tx.description or "",
            tx.doctor or "",
        ]


def monthly_xlsx(report: MonthlyReport) -> bytes:
    wb = Workbook()
    m = report.month
    sheets = t("export.sheets")
    headers = t("export.headers")
    s = t("export.summary")
    period = t(
        "export.summary.periodValue",
        label=m.label,
        start=to_jalali_str(m.start),
        end=to_jalali_str(m.end),
    )
    _sheet(
        wb,
        sheets["summary"],
        s["headers"],
        [
            (s["period"], period),
            (s["opening"], report.opening_receivable),
            (s["charged"], report.charged),
            (s["paid"], report.paid),
            (s["discounted"], report.discounted),
            (s["refunded"], report.refunded),
            (s["closing"], report.closing_receivable),
            (s["credit"], report.closing_credit),
            (s["debtors"], report.debtors_count),
        ],
        money_cols=(2,),
    )
    _sheet(
        wb,
        sheets["statement"],
        headers["statement"],
        (
            (
                r.name,
                r.opening,
                r.charged,
                r.paid,
                r.discounted,
                r.refunded,
                r.closing,
                r.pending_cheques,
            )
            for r in report.rows
        ),
        money_cols=range(2, 9),
    )
    _sheet(
        wb,
        sheets["byAccount"],
        headers["byAccount"],
        ((r.label, r.amount, r.count) for r in report.by_account),
        money_cols=(2,),
    )
    _sheet(
        wb,
        sheets["byMethod"],
        headers["byMethod"],
        ((r.label, r.amount, r.count) for r in report.by_method),
        money_cols=(2,),
    )
    _sheet(
        wb,
        sheets["chequesDue"],
        headers["chequesDue"],
        (
            (
                c.patient_name,
                c.amount,
                to_jalali_str(c.cheque_due_on),
                c.payer_name or t("common.self"),
                c.account_title or "",
                _label("cheque", c.cheque_status),
            )
            for c in report.cheques_due
        ),
        money_cols=(2,),
    )
    return _to_bytes(wb)


def full_backup_xlsx(db: Session) -> bytes:
    wb = Workbook()
    sheets = t("export.sheets")
    headers = t("export.headers")
    yes = t("common.yes")
    totals = ledger.patient_totals(db)
    patients = db.scalars(select(Patient).order_by(Patient.name)).all()
    _sheet(
        wb,
        sheets["patients"],
        headers["patients"],
        (
            (
                p.id,
                p.name,
                p.phone or "",
                totals[p.id].balance if p.id in totals else 0,
                totals[p.id].pending_cheques if p.id in totals else 0,
                yes if p.is_archived else "",
                p.note or "",
            )
            for p in patients
        ),
        money_cols=(4, 5),
    )
    txs = db.scalars(tx_select().order_by(Transaction.occurred_on, Transaction.id)).all()
    _sheet(wb, sheets["transactions"], headers["transactions"], _tx_rows(txs), money_cols=(6,))
    treatments = db.scalars(select(Treatment).order_by(Treatment.patient_id, Treatment.id)).all()
    pnames = {p.id: p.name for p in patients}
    _sheet(
        wb,
        sheets["cases"],
        headers["cases"],
        (
            (
                tr.id,
                pnames.get(tr.patient_id, ""),
                tr.title,
                _label("treatment", tr.kind),
                tr.doctor or "",
                _label("treatmentStatus", tr.status),
                tr.note or "",
            )
            for tr in treatments
        ),
    )
    accounts = db.scalars(select(Account).order_by(Account.id)).all()
    _sheet(
        wb,
        sheets["accounts"],
        headers["accounts"],
        (
            (
                a.id,
                a.title,
                a.owner_name or "",
                a.bank or "",
                _label("accountKind", a.kind),
                a.card_last4 or "",
                yes if a.is_archived else "",
            )
            for a in accounts
        ),
    )
    return _to_bytes(wb)


def month_sheet_xlsx(sheet: MonthSheet) -> bytes:
    """The month laid out like the practice's own sheet, under the headers the importer reads."""
    wb = Workbook()
    h = t("spreadsheet.headers")
    fields = ["name", "service", "balance", "payment", "date", "account", "doctor", "payer"]
    rows = (
        [
            r.patient_name,
            r.description or "",
            r.balance_after,
            r.amount if r.type == TxType.PAYMENT else 0,
            to_jalali_str(r.occurred_on),
            r.account_title or "",
            r.doctor or "",
            r.payer_name or "",
        ]
        for r in sheet.rows
    )
    _sheet(wb, sheet.month.label, [h[f][0] for f in fields], rows, money_cols=(3, 4))
    return _to_bytes(wb)
