"""A month sheet → standard workbook → database, end to end, with fake names."""

import io
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.core.db import new_session
from app.core.errors import ConflictError
from app.core.i18n import t
from app.models import ChequeStatus, PaymentMethod, TxType
from app.services.spreadsheet.cells import parse_amount, parse_date, parse_name
from app.services.spreadsheet.convert import convert_month
from app.services.spreadsheet.monthly import import_month
from app.services.workbook.importer import import_book
from app.services.workbook.reader import read_book
from app.services.workbook.writer import write_book
from app.utils.jalali import JMonth, parse_jalali
from tests.conftest import Api

M = 1_000_000
MEHR = (1405, 7)
ABAN = JMonth(1405, 8)


def _header(*, payer: bool = False) -> list[str]:
    """The sheet's header row, in the practice's own words (first spelling of each)."""
    h = t("spreadsheet.headers")
    fields = ["name", "service", "balance", "payment", "date", "account", "notes"]
    return [h[f][0] for f in fields] + ([h["payer"][0]] if payer else [])


def _workbook(path: Path) -> Path:
    blu = t("spreadsheet.banks")[next(iter(t("spreadsheet.banks")))][1]  # one-letter Blu
    doctor, cheque = t("spreadsheet.doctor"), t("spreadsheet.cheque")
    deposit = t("spreadsheet.deposit")
    visit, toothbrush = t("spreadsheet.salesWords")[0], t("spreadsheet.salesWords")[1]
    implant = t("spreadsheet.caseSuffixes")["implant"][1]
    wb = Workbook()
    mehr = wb.active
    mehr.title = "mehr"
    mehr.append(_header())
    mehr.append(["Ali", None, 9000 * M, 0, "1405-07-01", None, None])  # an earlier month: ignored

    aban = wb.create_sheet("aban")
    aban.append(_header(payer=True))
    aban.append(["Ali", None, 1000 * M, 0, "1405-08-01", None, None])  # owed before Aban
    aban.append(["Ali", None, 700 * M, 300 * M, "1405-08-05", f"{blu} Zahra", None, "Mohammad"])
    aban.append(["Ali", implant, 1200 * M, 0, "1405-08-06", None, f"{doctor} Emami"])
    aban.append(["Sara", deposit, 0, 50 * M, "1405-08-07", "pos", None])
    aban.append(["Reza", visit, 0, 5 * M, "1405-8-8", f"pos{blu}", None])  # loose date, glued POS
    aban.append(["Mina", None, 400 * M, 0, "1405-08-09", None, None])
    aban.append(["Mina", None, 100 * M, 0, "1405-08-10", None, None])  # dropped with no payment
    aban.append(["Mina", toothbrush, 0, 5 * M, "1405-08-11", "pos", None])  # a sale, not a payoff
    aban.append(["Omid", None, 650 * M, 0, "1405-08-25", f"{cheque} Forouhar", None])  # schedule
    aban.append(["Omid", None, 100 * M, 50 * M, "1405-08-12", "Ghomi", None])
    out = path / "practice.xlsx"
    wb.save(out)
    return out


def _azar(path: Path, rows: list[list[object]]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(_header())
    for row in rows:
        ws.append(row)
    out = path / "azar.xlsx"
    wb.save(out)
    return out


def test_cells() -> None:
    assert parse_amount("150,000,00").ambiguous and parse_amount("150,000,00").value == 15_000_000
    assert parse_amount(9.0).tiny
    assert parse_amount("?").text == "?"
    assert parse_name("Sara. Rahimi").key == parse_name("Sara Rahimi").key
    fixed = parse_date("1405-08-17", MEHR, None)
    assert fixed.on == parse_jalali("1405-07-17") and fixed.problem


@pytest.fixture
def book(tmp_path: Path):  # type: ignore[no-untyped-def]
    book, month = convert_month(_workbook(tmp_path))
    assert month == ABAN  # the latest month sheet
    return book


def test_reads_only_one_month_sheet(book) -> None:  # type: ignore[no-untyped-def]
    assert all(line.source.startswith("aban!") for line in book.lines if line.source)
    assert all(line.case is None for line in book.lines)  # one balance per patient, as the sheet


def test_another_month_can_be_picked(tmp_path: Path) -> None:
    book, month = convert_month(_workbook(tmp_path), JMonth(*MEHR))
    assert month == JMonth(*MEHR) and [x.amount for x in book.lines] == [9000 * M]


def test_rebuilds_charges_and_payments(book) -> None:  # type: ignore[no-untyped-def]
    ali = [line for line in book.lines if line.patient == "Ali"]
    assert [(line.type, line.amount) for line in ali] == [
        (TxType.CHARGE, 1000 * M),
        (TxType.PAYMENT, 300 * M),
        (TxType.CHARGE, 500 * M),
    ]
    assert ali[0].description == t("spreadsheet.openingNote")  # what was owed before the month
    assert ali[1].account and "Zahra" in ali[1].account
    assert ali[1].payer == "Mohammad"  # the payer column
    implant = ali[2]
    assert implant.description == t("spreadsheet.caseSuffixes")["implant"][1]  # the service
    assert implant.doctor == f"{t('spreadsheet.doctor')} emami"

    reza_pay = next(x for x in book.lines if x.patient == "Reza" and x.type == TxType.PAYMENT)
    assert reza_pay.method == PaymentMethod.POS
    assert reza_pay.account == t("spreadsheet.posTitle", bank=next(iter(t("spreadsheet.banks"))))

    patients = {p.name: p.balance for p in book.patients}
    assert patients["Ali"] == 1200 * M
    assert patients["Sara"] == -50 * M  # a deposit is credit
    assert patients["Reza"] == 0  # visit fee paid on the spot
    assert patients["Mina"] == 100 * M  # the toothbrush did not pay off her debt
    assert patients["Omid"] == 100 * M

    issues = " | ".join(i.text for i in book.issues)
    assert t("spreadsheet.issues.chequeSchedule") in issues


def test_round_trip_and_import(tmp_path: Path, book, api: Api) -> None:  # type: ignore[no-untyped-def]
    path = tmp_path / "standard.xlsx"
    path.write_bytes(write_book(book))
    result = read_book(path)
    assert result.errors == []
    assert len(result.book.lines) == len(book.lines)

    with new_session() as db:
        dry = import_book(db, result.book, dry_run=True)
    assert api.c.get("/api/patients").json() == []  # a dry run saves nothing

    with new_session() as db:
        summary = import_book(db, result.book)
    assert summary == dry
    balances = {p["name"]: p["balance"] for p in api.c.get("/api/patients").json()}
    assert (balances["Ali"], balances["Sara"], balances["Mina"]) == (1200 * M, -50 * M, 100 * M)


def test_next_month_continues_from_the_app_balances(tmp_path: Path, api: Api) -> None:
    with new_session() as db:
        import_month(db, _workbook(tmp_path), replace=True)
    azar = _azar(
        tmp_path,
        [
            ["Ali", None, 1100 * M, 100 * M, "1405-09-02", "Ghomi", None],  # paid 100 of 1200
            ["Nima", None, 80 * M, 20 * M, "1405-09-03", None, None],  # new this month
        ],
    )
    with new_session() as db:
        result = import_month(db, azar)
    assert result.month == JMonth(1405, 9)
    ali = [x for x in result.book.lines if x.patient == "Ali"]
    assert [(x.type, x.amount) for x in ali] == [(TxType.PAYMENT, 100 * M)]  # no charge invented
    balances = {p["name"]: p["balance"] for p in api.c.get("/api/patients").json()}
    assert (balances["Ali"], balances["Nima"]) == (1100 * M, 80 * M)

    with new_session() as db, pytest.raises(ConflictError):  # the same month twice
        import_month(db, azar)


def test_merge_into_combines_patients(tmp_path: Path, book, api: Api) -> None:  # type: ignore[no-untyped-def]
    for p in book.patients:
        if p.name == "Omid":
            p.merge_into = "Mina"
    path = tmp_path / "standard.xlsx"
    path.write_bytes(write_book(book))
    with new_session() as db:
        import_book(db, read_book(path).book)
    names = {p["name"]: p["balance"] for p in api.c.get("/api/patients").json()}
    assert "Omid" not in names and names["Mina"] == 200 * M


def test_reader_rejects_bad_rows(tmp_path: Path, book) -> None:  # type: ignore[no-untyped-def]
    book.lines[0].amount = 0
    path = tmp_path / "standard.xlsx"
    path.write_bytes(write_book(book))
    errors = read_book(path).errors
    assert len(errors) == 1 and "amount" in errors[0]


def test_import_refuses_a_non_empty_ledger(book, api: Api) -> None:  # type: ignore[no-untyped-def]
    api.tx(api.patient(), "charge", 1)
    with new_session() as db, pytest.raises(ConflictError):
        import_book(db, book)


def test_replace_swaps_an_earlier_import(book, api: Api) -> None:  # type: ignore[no-untyped-def]
    api.tx(api.patient("Old Patient"), "charge", 7)
    with new_session() as db:
        import_book(db, book, replace=True)
    names = {p["name"] for p in api.c.get("/api/patients").json()}
    assert "Old Patient" not in names and "Ali" in names
    rows = api.c.get("/api/transactions?limit=500").json()["items"]
    assert any(tx["doctor"] == f"{t('spreadsheet.doctor')} emami" for tx in rows)


# ── Regressions from earlier reviews ──────────────────────────────────────────


def test_repaired_dates_are_always_flagged_and_months_never_guessed() -> None:
    assert parse_date("1705-07-08", MEHR, None).problem  # century typo: fixed, but reported
    assert parse_date("1405-13-01", None, None).on is None  # no month 13 → no guess
    assert parse_date("1405-07-08", MEHR, None).problem is None


def test_negative_amounts_are_flagged() -> None:
    assert parse_amount(-5 * M).negative and parse_amount(-5 * M).value == 5 * M


def test_keywords_match_whole_words_not_name_prefixes() -> None:
    from app.services.spreadsheet.destinations import Accounts

    persian_pos = next(w for w in t("spreadsheet.posWords") if not w.isascii())
    jaw = next(w for w in t("spreadsheet.noteWords") if len(w) == 2)
    accounts = Accounts()
    physician = accounts.read(persian_pos + "\u0634\u06a9")  # a word that merely starts with it
    assert physician.method is None and physician.account
    surname = accounts.read(jaw + "\u0631\u06cc")  # a surname starting with the note word
    assert surname.account and not surname.notes
    glued = accounts.read("pos" + t("spreadsheet.banks")[next(iter(t("spreadsheet.banks")))][0])
    assert glued.method == PaymentMethod.POS and glued.account != t("spreadsheet.posTitleNoBank")


def test_two_sheets_of_the_same_month_do_not_crash(tmp_path: Path) -> None:
    wb = Workbook()
    for title in ("a", "b"):
        ws = wb.create_sheet(title)
        ws.append(_header())
        ws.append(["Ali", None, 100 * M, 0, "1405-07-01", None, None])
    path = tmp_path / "dup.xlsx"
    wb.save(path)
    assert convert_month(path)[0].lines


def test_cheques_assumed_cleared_are_flagged(tmp_path: Path) -> None:
    path = _azar(tmp_path, [["Ali", None, 0, 100 * M, "1405-09-01", t("spreadsheet.cheque"), None]])
    (line,) = [x for x in convert_month(path)[0].lines if x.type == TxType.PAYMENT]
    assert line.cheque_status == ChequeStatus.CLEARED
    assert line.review and t("spreadsheet.issues.chequeAssumedCleared") in line.review


def test_reader_rejects_negative_and_overlong_values(tmp_path: Path, book) -> None:  # type: ignore[no-untyped-def]
    book.lines[0].payer = "x" * 121
    path = tmp_path / "standard.xlsx"
    path.write_bytes(write_book(book))
    from openpyxl import load_workbook

    wb = load_workbook(path)
    wb[t("workbook.sheets")["transactions"]].cell(row=3, column=5).value = -5
    wb.save(path)
    errors = read_book(path).errors
    assert len(errors) == 2
    assert "payer too long" in errors[0] and "amount" in errors[1]


def test_force_import_reuses_existing_patients_and_accounts(tmp_path: Path, book, api: Api) -> None:  # type: ignore[no-untyped-def]
    with new_session() as db:
        import_book(db, book)
    first = {p["name"] for p in api.c.get("/api/patients").json()}
    accounts = len(api.c.get("/api/accounts").json())
    with new_session() as db:
        import_book(db, book, force=True)
    patients = [p["name"] for p in api.c.get("/api/patients").json()]
    assert sorted(patients) == sorted(first)  # no duplicates
    assert len(api.c.get("/api/accounts").json()) == accounts


def test_cli_imports_without_production_settings() -> None:
    """`hash-password` must work before AUTH_PASSWORD_HASH exists (first install)."""
    import os
    import subprocess
    import sys

    env = {k: v for k, v in os.environ.items() if k not in ("AUTH_PASSWORD_HASH", "SECRET_KEY")}
    env.update(ENV="production", DATABASE_URL="sqlite:////nonexistent/taraz.db")
    cli_only = subprocess.run(
        [sys.executable, "-c", "import app.cli"],
        env=env,
        capture_output=True,
        text=True,
        cwd=Path(__file__).resolve().parent.parent,
    )
    assert cli_only.returncode == 0, cli_only.stderr


def test_cli_imports_a_month_from_stdin(tmp_path: Path, api: Api, monkeypatch, capsys) -> None:  # type: ignore[no-untyped-def]
    """The production container is read-only, so the file is piped in, not copied."""
    import sys

    from app.cli import main

    stdin = io.TextIOWrapper(io.BytesIO(_workbook(tmp_path).read_bytes()))
    monkeypatch.setattr(sys, "stdin", stdin)
    assert main(["import-month", "-", "--dry-run"]) == 0
    assert "[dry run, nothing saved]" in capsys.readouterr().out
    assert api.c.get("/api/patients").json() == []
