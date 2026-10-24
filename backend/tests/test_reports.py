from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from app.core.i18n import t
from app.utils.jalali import JMonth, parse_jalali
from tests.conftest import Api

# Mehr 1405 = 2026-09-23 … 2026-10-22; Aban 1405 starts 2026-10-23.
MEHR = JMonth(1405, 7)


def test_jalali_month_boundaries() -> None:
    assert (MEHR.start.isoformat(), MEHR.end.isoformat()) == ("2026-09-23", "2026-10-22")
    assert JMonth(1405, 12).next() == JMonth(1406, 1)
    assert parse_jalali("1405,04,30") == parse_jalali("1405/4/30")
    assert parse_jalali("1405-067-07") is None


def test_month_end_statement(api: Api) -> None:
    pid = api.patient("Monthly Patient")
    acc = api.account("Blu", color="sky")
    api.tx(pid, "charge", 1_000, on="2026-09-10")  # Shahrivar → opening balance
    api.tx(pid, "charge", 500, on="2026-10-01")
    api.tx(pid, "payment", 300, on="2026-10-05", account_id=acc, method="card")
    api.tx(pid, "payment", 100, on="2026-10-23", method="cash")  # Aban — outside the month
    # Received in Mehr but due in Aban: still in hand at month end.
    api.tx(pid, "payment", 200, on="2026-10-10", method="cheque", cheque_due_on="2026-11-05")

    r = api.c.get("/api/reports/monthly", params={"year": 1405, "month": 7}).json()
    assert r["month"]["label"] == f"{t('months')[6]} 1405"
    (row,) = r["rows"]
    assert (row["opening"], row["charged"], row["paid"], row["closing"]) == (1_000, 500, 300, 1_200)
    assert row["pending_cheques"] == 200
    assert (r["opening_receivable"], r["closing_receivable"], r["debtors_count"]) == (
        1_000,
        1_200,
        1,
    )
    assert [(a["label"], a["amount"]) for a in r["by_account"]] == [("Blu", 300)]
    assert r["cheques_due"] == []


def test_cleared_cheque_lands_on_its_due_date(api: Api) -> None:
    pid = api.patient()
    api.tx(pid, "charge", 1_000, on="2026-09-25")
    cheque = api.tx(
        pid, "payment", 400, on="2026-09-30", method="cheque", cheque_due_on="2026-10-28"
    )
    api.c.patch(f"/api/transactions/{cheque['id']}", json={"cheque_status": "cleared"})

    mehr = api.c.get("/api/reports/monthly", params={"year": 1405, "month": 7}).json()
    aban = api.c.get("/api/reports/monthly", params={"year": 1405, "month": 8}).json()
    assert mehr["rows"][0]["closing"] == 1_000 and mehr["paid"] == 0
    assert aban["rows"][0]["closing"] == 600 and aban["paid"] == 400
    assert [c["id"] for c in aban["cheques_due"]] == [cheque["id"]]


def test_monthly_report_rejects_bad_month(api: Api) -> None:
    assert api.c.get("/api/reports/monthly", params={"year": 1405, "month": 13}).status_code == 422


def test_dashboard(api: Api) -> None:
    a, b = api.patient("A"), api.patient("B")
    api.tx(a, "charge", 700)
    api.tx(b, "payment", 50, method="cash")
    api.tx(a, "payment", 100, method="cheque", cheque_due_on="2026-12-01")
    d = api.c.get("/api/dashboard").json()
    assert (d["receivable"], d["credit"], d["debtors_count"]) == (700, 50, 1)
    assert (d["pending_cheques_total"], d["pending_cheques_count"]) == (100, 1)
    assert len(d["series"]) == 6
    assert d["top_debtors"][0]["name"] == "A"


def test_excel_exports(api: Api) -> None:
    pid = api.patient()
    api.tx(pid, "charge", 1_000, on="2026-10-01")
    r = api.c.get("/api/reports/monthly.xlsx", params={"year": 1405, "month": 7})
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    sheets = t("export.sheets")
    assert wb.sheetnames[:2] == [sheets["summary"], sheets["statement"]]
    assert wb[sheets["statement"]].sheet_view.rightToLeft

    backup = load_workbook(BytesIO(api.c.get("/api/export/backup.xlsx").content))
    assert backup[sheets["transactions"]].max_row == 2


def test_sqlite_backup_is_readable_and_pruned(tmp_path: Path, api: Api) -> None:
    from app.core.config import get_settings
    from app.services.backup import backup_sqlite

    api.patient("Backup Check")
    db_path = Path(get_settings().database_url.removeprefix("sqlite:///"))
    for _ in range(3):
        last = backup_sqlite(db_path, tmp_path, keep=2)
    import sqlite3

    assert len(list(tmp_path.glob("taraz-*.db"))) <= 2
    with sqlite3.connect(last) as conn:
        assert conn.execute("select name from patients").fetchall() == [("Backup Check",)]
