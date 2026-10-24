"""Old workbook → standard workbook → database, end to end, with fake names."""

from pathlib import Path

import pytest
from openpyxl import Workbook

from app.core.db import new_session
from app.core.i18n import t
from app.models import ChequeStatus, PaymentMethod, TreatmentKind, TxType
from app.services.legacy.cells import parse_amount, parse_date, parse_name
from app.services.legacy.convert import convert_legacy
from app.services.workbook.importer import import_book
from app.services.workbook.reader import read_book
from app.services.workbook.writer import write_book
from app.utils.jalali import parse_jalali
from tests.conftest import Api

M = 1_000_000
AS_OF = parse_jalali("1405-08-01")
MEHR = (1405, 7)


def _legacy_workbook(path: Path) -> Path:
    blu = t("legacy.banks")[next(iter(t("legacy.banks")))][1]  # the one-letter spelling of Blu
    doctor, cheque, deposit = t("legacy.doctor"), t("legacy.cheque"), t("legacy.deposit")
    visit = t("legacy.salesWords")[0]
    wb = Workbook()
    month = wb.active
    month.title = t("legacy.monthSheets")[3]
    month.append(["name", "balance", "paid", "date", "account", "note"])
    month.append(["Ali", 1000 * M, 0, "1405-07-01", None, None])  # treatment worth 1000
    month.append(["Ali", 700 * M, 300 * M, "1405-07-05", f"{blu} Zahra", None])
    month.append(["Ali imp", 0, 500 * M, "1405-07-06", f"{doctor} Emami", f"{blu} Zahra"])
    month.append(["Sara", deposit, 50 * M, "1405-07-07", "pos", None])
    month.append(["Reza", visit, 5 * M, "1405-7-8", f"pos{blu}", None])  # loose date, glued POS
    month.append(["Mina", 400 * M, 0, "1405-07-09", None, None])
    month.append(["Mina", 100 * M, 0, "1405-07-10", None, None])  # dropped with no payment
    month.append(
        ["Omid", 650 * M, 0, "1405-09-25", f"{cheque} Forouhar", None]
    )  # a cheque schedule
    month.append(["Omid", 100 * M, 50 * M, "1405-08-02", "Ghomi", None])  # Aban date in Mehr sheet

    debtors = wb.create_sheet(t("legacy.debtorsSheet"))
    debtors.append(["name", "amount", "date", "account", "note"])
    debtors.append(["Ali", 700 * M, None, None, None])
    debtors.append(["Ali", 300 * M, "1405-08-10", f"{cheque} Forouhar", None])
    debtors.append(["Newcomer", 200 * M, None, None, None])
    out = path / "old.xlsx"
    wb.save(out)
    return out


def test_cells() -> None:
    assert parse_amount("150,000,00").ambiguous and parse_amount("150,000,00").value == 15_000_000
    assert parse_amount(9.0).tiny
    assert parse_amount("?").text == "?"
    assert parse_name("Sara Rahimiimp").case == TreatmentKind.IMPLANT  # suffix glued to the name
    assert parse_name("Sara. Rahimi").key == parse_name("Sara Rahimi").key
    fixed = parse_date("1405-08-17", MEHR, None)
    assert fixed.on == parse_jalali("1405-07-17") and fixed.problem


@pytest.fixture
def book(tmp_path: Path):  # type: ignore[no-untyped-def]
    return convert_legacy(_legacy_workbook(tmp_path), AS_OF)


def test_rebuilds_charges_payments_and_debtors(book) -> None:  # type: ignore[no-untyped-def]
    ali = [line for line in book.lines if line.patient == "Ali" and line.case is None]
    assert [(line.type, line.amount) for line in ali[:3]] == [
        (TxType.CHARGE, 1000 * M),
        (TxType.PAYMENT, 300 * M),
        (TxType.CHARGE, 300 * M),  # debtors list adds the cheque-backed 300
    ]
    assert ali[1].account and "Zahra" in ali[1].account
    cheque = ali[3]
    assert (cheque.method, cheque.cheque_status, cheque.payer) == (
        PaymentMethod.CHEQUE,
        ChequeStatus.PENDING,
        "Forouhar",
    )

    reza_pay = next(x for x in book.lines if x.patient == "Reza" and x.type == TxType.PAYMENT)
    assert reza_pay.method == PaymentMethod.POS
    assert reza_pay.account == t("legacy.posTitle", bank=next(iter(t("legacy.banks"))))

    implant = [line for line in book.lines if line.case == TreatmentKind.IMPLANT]
    assert {line.doctor for line in implant} == {f"{t('legacy.doctor')} emami"}

    patients = {p.name: p for p in book.patients}
    assert (patients["Ali"].balance, patients["Ali"].pending_cheques) == (1000 * M, 300 * M)
    assert patients["Sara"].balance == -50 * M  # a deposit is credit
    assert patients["Reza"].balance == 0  # visit fee paid on the spot
    assert patients["Mina"].balance == 100 * M
    assert patients["Newcomer"].balance == 200 * M

    issues = " | ".join(i.text for i in book.issues)
    assert t("legacy.issues.chequeSchedule") in issues
    assert "Omid" in {i.patient for i in book.issues}


def test_round_trip_and_import(tmp_path: Path, book, api: Api) -> None:  # type: ignore[no-untyped-def]
    path = tmp_path / "standard.xlsx"
    path.write_bytes(write_book(book))
    result = read_book(path)
    assert result.errors == []
    assert len(result.book.lines) == len(book.lines)

    with new_session() as db:
        dry = import_book(db, result.book, dry_run=True)
    assert api.c.get("/api/patients").json() == []  # a dry run saves nothing

    with new_session() as db:
        summary = import_book(db, result.book)
    assert summary == dry
    balances = {p["name"]: p for p in api.c.get("/api/patients").json()}
    assert balances["Ali"]["balance"] == 1000 * M
    assert balances["Ali"]["pending_cheques"] == 300 * M
    assert balances["Sara"]["balance"] == -50 * M


def test_merge_into_combines_patients(tmp_path: Path, book, api: Api) -> None:  # type: ignore[no-untyped-def]
    for p in book.patients:
        if p.name == "Newcomer":
            p.merge_into = "Mina"
    path = tmp_path / "standard.xlsx"
    path.write_bytes(write_book(book))
    with new_session() as db:
        import_book(db, read_book(path).book)
    names = {p["name"]: p["balance"] for p in api.c.get("/api/patients").json()}
    assert "Newcomer" not in names and names["Mina"] == 300 * M


def test_reader_rejects_bad_rows(tmp_path: Path, book) -> None:  # type: ignore[no-untyped-def]
    book.lines[0].amount = 0
    path = tmp_path / "standard.xlsx"
    path.write_bytes(write_book(book))
    errors = read_book(path).errors
    assert len(errors) == 1 and "amount" in errors[0]


def test_import_refuses_a_non_empty_ledger(book, api: Api) -> None:  # type: ignore[no-untyped-def]
    from app.core.errors import ConflictError

    api.tx(api.patient(), "charge", 1)
    with new_session() as db, pytest.raises(ConflictError):
        import_book(db, book)


def test_settle_unlisted_closes_balances_missing_from_the_debtors_list(tmp_path: Path) -> None:
    kept = {p.name: p.balance for p in convert_legacy(_legacy_workbook(tmp_path), AS_OF).patients}
    settled = convert_legacy(_legacy_workbook(tmp_path), AS_OF, settle_unlisted=True)
    balances = {p.name: p.balance for p in settled.patients}
    assert kept["Mina"] == 100 * M and balances["Mina"] == 0  # not on the debtors list
    assert balances["Ali"] == kept["Ali"]  # on the list: untouched


# ── Regressions from the code review ─────────────────────────────────────────


def test_repaired_dates_are_always_flagged_and_months_never_guessed() -> None:
    assert parse_date("1705-07-08", MEHR, None).problem  # century typo: fixed, but reported
    assert parse_date("1405-13-01", None, None).on is None  # no month 13 → no guess
    assert parse_date("1405-07-08", MEHR, None).problem is None


def test_negative_amounts_are_flagged() -> None:
    assert parse_amount(-5 * M).negative and parse_amount(-5 * M).value == 5 * M


def test_keywords_match_whole_words_not_name_prefixes() -> None:
    from app.services.legacy.destinations import Accounts

    persian_pos = next(w for w in t("legacy.posWords") if not w.isascii())
    jaw = next(w for w in t("legacy.noteWords") if len(w) == 2)
    accounts = Accounts()
    physician = accounts.read(persian_pos + "\u0634\u06a9")  # a word that merely starts with it
    assert physician.method is None and physician.account
    surname = accounts.read(jaw + "\u0631\u06cc")  # a surname starting with the note word
    assert surname.account and not surname.notes
    glued = accounts.read("pos" + t("legacy.banks")[next(iter(t("legacy.banks")))][0])
    assert glued.method == PaymentMethod.POS and glued.account != t("legacy.posTitleNoBank")


def test_two_sheets_of_the_same_month_do_not_crash(tmp_path: Path) -> None:
    wb = Workbook()
    for title in t("legacy.monthSheets")[3:5]:
        ws = wb.create_sheet(title)
        ws.append(["name", "balance", "paid", "date", "account", "note"])
        ws.append(["Ali", 100 * M, 0, "1405-07-01", None, None])
    path = tmp_path / "dup.xlsx"
    wb.save(path)
    assert convert_legacy(path, AS_OF).lines


def test_cheques_assumed_cleared_are_flagged(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = t("legacy.monthSheets")[3]
    ws.append(["name", "balance", "paid", "date", "account", "note"])
    ws.append(["Ali", 0, 100 * M, "1405-07-01", t("legacy.cheque"), None])
    path = tmp_path / "cheque.xlsx"
    wb.save(path)
    (line,) = [x for x in convert_legacy(path, AS_OF).lines if x.type == TxType.PAYMENT]
    assert line.cheque_status == ChequeStatus.CLEARED
    assert line.review and t("legacy.issues.chequeAssumedCleared") in line.review


def test_reader_rejects_negative_and_overlong_values(tmp_path: Path, book) -> None:  # type: ignore[no-untyped-def]
    book.lines[0].payer = "x" * 121
    path = tmp_path / "standard.xlsx"
    path.write_bytes(write_book(book))
    from openpyxl import load_workbook

    wb = load_workbook(path)
    wb[t("workbook.sheets")["transactions"]].cell(row=3, column=5).value = -5
    wb.save(path)
    errors = read_book(path).errors
    assert len(errors) == 2
    assert "payer too long" in errors[0] and "amount" in errors[1]


def test_force_import_reuses_existing_patients_and_accounts(tmp_path: Path, book, api: Api) -> None:  # type: ignore[no-untyped-def]
    with new_session() as db:
        import_book(db, book)
    first = {p["name"] for p in api.c.get("/api/patients").json()}
    accounts = len(api.c.get("/api/accounts").json())
    with new_session() as db:
        import_book(db, book, force=True)
    patients = [p["name"] for p in api.c.get("/api/patients").json()]
    assert sorted(patients) == sorted(first)  # no duplicates
    assert len(api.c.get("/api/accounts").json()) == accounts


def test_cli_imports_without_production_settings() -> None:
    """`hash-password` must work before AUTH_PASSWORD_HASH exists (first install)."""
    import os
    import subprocess
    import sys

    env = {k: v for k, v in os.environ.items() if k not in ("AUTH_PASSWORD_HASH", "SECRET_KEY")}
    env.update(ENV="production", DATABASE_URL="sqlite:////nonexistent/taraz.db")
    cli_only = subprocess.run(
        [sys.executable, "-c", "import app.cli"],
        env=env,
        capture_output=True,
        text=True,
        cwd=Path(__file__).resolve().parent.parent,
    )
    assert cli_only.returncode == 0, cli_only.stderr


def test_cli_reads_the_workbook_from_stdin(book, api: Api, monkeypatch, capsys) -> None:  # type: ignore[no-untyped-def]
    """The production container is read-only, so the file is piped in, not copied."""
    import io
    import sys

    from app.cli import main

    stdin = io.TextIOWrapper(io.BytesIO(write_book(book)))
    monkeypatch.setattr(sys, "stdin", stdin)
    assert main(["import-workbook", "-", "--dry-run"]) == 0
    assert "[dry run, nothing saved]" in capsys.readouterr().out
    assert api.c.get("/api/patients").json() == []
